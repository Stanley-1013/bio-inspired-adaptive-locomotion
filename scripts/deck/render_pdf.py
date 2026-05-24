"""Faithful-enough PDF renderer for a pptx (LibreOffice unavailable in sandbox).
Reads real shape geometry, fills, lines, and per-run text style from the file."""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

DPI = 150
EMU = 914400
PX = DPI  # px per inch

REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
_fc = {}
def font(sz, bold):
    sz = max(8, int(round(sz)))
    key = (sz, bold)
    if key not in _fc:
        _fc[key] = ImageFont.truetype(BLD if bold else REG, sz)
    return _fc[key]

def emu_px(v):
    return None if v is None else int(round(v / EMU * PX))

def rgb(c):
    return (c[0], c[1], c[2])

def shape_fill(sh):
    try:
        if sh.fill.type == 1:  # solid
            return rgb(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None

def line_rgb(sh):
    try:
        if sh.line.color and sh.line.color.type is not None:
            return rgb(sh.line.color.rgb)
    except Exception:
        pass
    return None

def line_w(sh):
    try:
        if sh.line.width:
            return max(1, int(round(sh.line.width / EMU * PX)))
    except Exception:
        pass
    return 1

def run_color(r, default=(36, 50, 48)):
    try:
        if r.font.color and r.font.color.type is not None:
            return rgb(r.font.color.rgb)
    except Exception:
        pass
    return default

def run_size(r, default=14):
    try:
        if r.font.size:
            return r.font.size.pt
    except Exception:
        pass
    return default

def para_tokens(p):
    toks = []
    for r in p.runs:
        col = run_color(r)
        sz = run_size(r) * DPI / 72.0
        bold = bool(r.font.bold)
        words = r.text.split(" ")
        for j, w in enumerate(words):
            toks.append({"t": w, "c": col, "s": sz, "b": bold,
                         "sp": (j < len(words) - 1)})  # trailing space
    return toks

def draw_text(d, sh, l, t, w, h):
    tf = sh.text_frame
    ml = emu_px(tf.margin_left) or 0
    mr = emu_px(tf.margin_right) or 0
    mt = emu_px(tf.margin_top) or 0
    mb = emu_px(tf.margin_bottom) or 0
    aw = max(10, w - ml - mr)
    x0 = l + ml
    # wrap each paragraph into lines
    lines = []  # each: (tokens, height)
    for p in tf.paragraphs:
        toks = para_tokens(p)
        align = p.alignment
        if not toks:
            sz = 14 * DPI / 72.0
            lines.append(([], int(sz * 1.25), align))
            continue
        cur, curw = [], 0
        for tk in toks:
            f = font(tk["s"], tk["b"])
            wpx = d.textlength(tk["t"], font=f)
            spw = d.textlength(" ", font=f) if tk["sp"] else 0
            if cur and curw + wpx > aw:
                lines.append((cur, max(int(x["s"] * 1.28) for x in cur), align))
                cur, curw = [], 0
            cur.append(tk)
            curw += wpx + spw
        if cur:
            lines.append((cur, max(int(x["s"] * 1.28) for x in cur), align))
    total = sum(lh for _, lh, _ in lines)
    va = tf.vertical_anchor
    if va == MSO_ANCHOR.MIDDLE:
        y = t + mt + max(0, (h - mt - mb - total) // 2)
    elif va == MSO_ANCHOR.BOTTOM:
        y = t + h - mb - total
    else:
        y = t + mt
    for toks, lh, align in lines:
        lw = 0
        for tk in toks:
            f = font(tk["s"], tk["b"])
            lw += d.textlength(tk["t"], font=f) + (d.textlength(" ", font=f) if tk["sp"] else 0)
        if align == PP_ALIGN.CENTER:
            x = x0 + max(0, (aw - lw) // 2)
        elif align == PP_ALIGN.RIGHT:
            x = x0 + max(0, aw - lw)
        else:
            x = x0
        for tk in toks:
            f = font(tk["s"], tk["b"])
            d.text((x, y), tk["t"], fill=tk["c"], font=f)
            x += d.textlength(tk["t"], font=f) + (d.textlength(" ", font=f) if tk["sp"] else 0)
        y += lh

prs = Presentation(sys.argv[1])
W = emu_px(prs.slide_width)
H = emu_px(prs.slide_height)
pages = []
for slide in prs.slides:
    try:
        bg = rgb(slide.background.fill.fore_color.rgb)
    except Exception:
        bg = (244, 241, 234)
    img = Image.new("RGB", (W, H), bg)
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        l, t = emu_px(sh.left), emu_px(sh.top)
        w, h = emu_px(sh.width), emu_px(sh.height)
        if l is None or t is None:
            continue
        st = sh.shape_type
        name = str(st)
        is_line = (st == MSO_SHAPE_TYPE.LINE) or ("CONNECTOR" in name) or ("LINE" in name)
        is_oval = False
        if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                is_oval = "OVAL" in str(sh.auto_shape_type)
            except Exception:
                is_oval = False
        if is_line:
            lc = line_rgb(sh) or (200, 150, 60)
            d.line([(l, t), (l + (w or 0), t + (h or 0))], fill=lc, width=line_w(sh))
            continue
        fill = shape_fill(sh)
        lc = line_rgb(sh)
        if w and h and (fill is not None or lc is not None):
            box = [l, t, l + w, t + h]
            if is_oval:
                d.ellipse(box, fill=fill, outline=lc, width=line_w(sh) if lc else 1)
            else:
                d.rectangle(box, fill=fill, outline=lc, width=line_w(sh) if lc else 1)
        if sh.has_text_frame and sh.text_frame.text.strip():
            draw_text(d, sh, l, t, w or 0, h or 0)
    pages.append(img)

out = sys.argv[2]
pages[0].save(out, "PDF", resolution=float(DPI), save_all=True, append_images=pages[1:])
print("wrote", out, "pages:", len(pages))
